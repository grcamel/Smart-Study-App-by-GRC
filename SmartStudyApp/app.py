from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, session, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
import google.generativeai as genai
import PyPDF2
import pptx
import docx
import json
import os
import re
import io
import markdown
import sqlite3
import time
from datetime import datetime, date, timedelta
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

app = Flask(__name__)
app.secret_key = "materiku_rahasia_pplg_2026"
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max


genai.configure(api_key="AQ.Ab8RN6IpGtVIVZQMIgHSG1Xgp9RQGnU5UUwLhhG6hfvMONMorA")
model = genai.GenerativeModel('gemini-flash-latest')

# --- DATABASE SQLITE ---
def init_db():
    conn = sqlite3.connect('materiku.db')
    cursor = conn.cursor()

    cursor.execute('''CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT, nama TEXT NOT NULL,
        email TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
        tanggal_daftar DATETIME DEFAULT CURRENT_TIMESTAMP)''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS materi (
        id INTEGER PRIMARY KEY AUTOINCREMENT, judul TEXT NOT NULL,
        jenjang TEXT DEFAULT 'SMA', tipe_input TEXT, isi_teks TEXT NOT NULL,
        tanggal DATETIME DEFAULT CURRENT_TIMESTAMP, user_id INTEGER)''')

    # Migrasi aman: kalau tabel 'materi' sudah ada dari versi sebelumnya (tanpa kolom user_id),
    # tambahkan kolomnya. Kalau kolomnya sudah ada, abaikan error-nya.
    try:
        cursor.execute('ALTER TABLE materi ADD COLUMN user_id INTEGER')
    except sqlite3.OperationalError:
        pass  # kolom sudah ada

    cursor.execute('''CREATE TABLE IF NOT EXISTS hasil_belajar (
        id INTEGER PRIMARY KEY AUTOINCREMENT, materi_id INTEGER,
        jenis_hasil TEXT, konten TEXT,
        FOREIGN KEY (materi_id) REFERENCES materi(id))''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS kuis (
        id INTEGER PRIMARY KEY AUTOINCREMENT, materi_id INTEGER,
        soal TEXT, pilihan TEXT, kunci_jawaban TEXT, pembahasan TEXT,
        FOREIGN KEY (materi_id) REFERENCES materi(id))''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS progress (
        id INTEGER PRIMARY KEY AUTOINCREMENT, materi_id INTEGER,
        total_soal INTEGER, benar INTEGER, nilai REAL,
        waktu_kuis DATETIME, FOREIGN KEY (materi_id) REFERENCES materi(id))''')

    conn.commit()
    conn.close()

def get_db():
    conn = sqlite3.connect('materiku.db')
    conn.row_factory = sqlite3.Row
    return conn

# --- AUTH HELPER ---
def login_required(f):
    """Decorator: halaman ini cuma bisa diakses kalau sudah login."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            flash('Silakan login terlebih dahulu.', 'danger')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# --- FUNGSI EKSTRAK TEKS ---
def extract_text_from_file(filepath, filename):
    text = ""
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        reader = PyPDF2.PdfReader(filepath)
        for page in reader.pages: text += page.extract_text() + "\n"
    elif ext == '.pptx':
        prs = pptx.Presentation(filepath)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"): text += shape.text + "\n"
    elif ext == '.docx':
        doc = docx.Document(filepath)
        for paragraph in doc.paragraphs: text += paragraph.text + "\n"
    else:
        with open(filepath, 'r', encoding='utf-8') as f: text = f.read()
    return text.strip()

# --- FUNGSI AI ---
def clean_json(response_text):
    match = re.search(r'\[.*\]', response_text, re.DOTALL)
    if match: return json.loads(match.group(0))
    return json.loads(response_text)

def clean_json_object(response_text):
    """Ambil objek JSON {...} dari teks respons AI (kadang dibungkus ```json ... ```)."""
    match = re.search(r'\{.*\}', response_text, re.DOTALL)
    if match:
        return json.loads(match.group(0))
    return json.loads(response_text)

def markdown_ke_flowables(teks, styles):
    """Konversi teks markdown sederhana (heading, bold, italic, bullet) menjadi
    daftar elemen reportlab (Paragraph/Spacer) untuk dimasukkan ke PDF."""
    flowables = []
    if not teks or not teks.strip():
        return [Paragraph('-', styles['Normal'])]

    bullet_buffer = []
    def flush_bullets():
        for b in bullet_buffer:
            flowables.append(Paragraph(f"&bull;&nbsp; {b}", styles['Normal']))
        bullet_buffer.clear()

    for baris in teks.split('\n'):
        baris = baris.strip()
        if not baris:
            flush_bullets()
            flowables.append(Spacer(1, 6))
            continue

        # escape karakter HTML dasar, lalu terapkan bold/italic markdown
        aman = baris.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        aman = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', aman)
        aman = re.sub(r'\*(.+?)\*', r'<i>\1</i>', aman)

        if baris.startswith('### '):
            flush_bullets(); flowables.append(Paragraph(aman[4:], styles['Heading3']))
        elif baris.startswith('## '):
            flush_bullets(); flowables.append(Paragraph(aman[3:], styles['Heading2']))
        elif baris.startswith('# '):
            flush_bullets(); flowables.append(Paragraph(aman[2:], styles['Heading2']))
        elif baris.startswith('- ') or baris.startswith('* '):
            bullet_buffer.append(aman[2:])
        else:
            flush_bullets(); flowables.append(Paragraph(aman, styles['Normal']))
    flush_bullets()
    return flowables

def generate_paket_belajar(teks, jenjang='SMA'):
    # Semua bagian digabung jadi SATU prompt & SATU panggilan API,
    # supaya hemat kuota (dulu 6 request/materi, sekarang cuma 1 request/materi).
    prompt = f"""Kamu adalah asisten belajar untuk siswa {jenjang}. Berdasarkan materi di bawah, buatkan paket belajar lengkap.

BALAS HANYA DENGAN SATU OBJEK JSON (tanpa teks lain, tanpa markdown code fence), dengan struktur PERSIS seperti ini:

{{
  "ringkasan": "penjelasan materi dengan bahasa mudah dipahami, format Markdown (heading, bullet points)",
  "poin_penting": "5-8 poin paling penting, format Markdown bullet points",
  "flashcard": [{{"q": "pertanyaan", "a": "jawaban"}}, ... jumlahnya menyesuaikan kekayaan materi: buat SEBANYAK MUNGKIN konsep/fakta/istilah penting yang benar-benar ada di materi (jangan mengarang atau mengulang-ulang pertanyaan yang mirip hanya untuk menambah jumlah). Untuk materi singkat cukup sekitar 5-8 kartu, untuk materi yang kaya/panjang boleh sampai 15-25 kartu, minimal tetap 5],
  "kuis": [{{"soal": "...", "opsi": ["A. ...", "B. ...", "C. ...", "D. ..."], "kunci": "A", "pembahasan": "..."}}, ... total 20 item],
  "kesimpulan": "kesimpulan singkat, 2 paragraf",
  "rencana_belajar": "rencana belajar 15-30 menit, dibagi sesi-sesi kecil, format Markdown"
}}

MATERI:
{teks[:4000]}
"""

    kosong = {
        'ringkasan': "Gagal generate konten ini.",
        'poin_penting': "Gagal generate konten ini.",
        'flashcard': "[]",
        'kuis': "[]",
        'kesimpulan': "Gagal generate konten ini.",
        'rencana_belajar': "Gagal generate konten ini.",
    }

    kuota_habis = {
        'ringkasan': "**Kuota harian Gemini API (free tier) sudah habis.** Kuota gratis reset otomatis tiap hari (sekitar siang/sore WIB). Coba lagi nanti, atau aktifkan billing di Google AI Studio untuk kuota lebih besar.",
        'poin_penting': "Kuota harian API sudah habis, coba lagi nanti setelah reset.",
        'flashcard': "[]",
        'kuis': "[]",
        'kesimpulan': "Kuota harian API sudah habis, coba lagi nanti setelah reset.",
        'rencana_belajar': "Kuota harian API sudah habis, coba lagi nanti setelah reset.",
    }

    max_percobaan = 3
    for percobaan in range(1, max_percobaan + 1):
        try:
            print(f"[INFO] Generate paket belajar (percobaan {percobaan}/{max_percobaan})...")
            response = model.generate_content(
                prompt,
                request_options={"timeout": 180}
            )
            data = clean_json_object(response.text)

            hasil = {
                'ringkasan': data.get('ringkasan', ''),
                'poin_penting': data.get('poin_penting', ''),
                'flashcard': json.dumps(data.get('flashcard', [])),
                'kuis': json.dumps(data.get('kuis', [])),
                'kesimpulan': data.get('kesimpulan', ''),
                'rencana_belajar': data.get('rencana_belajar', ''),
            }
            print("[INFO] Paket belajar berhasil dibuat dalam 1 kali panggilan API.")
            return hasil
        except Exception as e:
            error_text = str(e)
            print(f"[ERROR] Gagal generate paket belajar (percobaan {percobaan}): {error_text}")

            # Kalau ini kuota HARIAN yang habis, retry percuma (baru reset besok) -> langsung stop.
            if "PerDay" in error_text or "GenerateRequestsPerDayPerProjectPerModel" in error_text:
                print("[INFO] Kuota HARIAN habis, tidak perlu retry (baru reset di hari berikutnya).")
                return kuota_habis

            if percobaan == max_percobaan:
                return kosong
            # Beri jeda sebelum retry supaya tidak nabrak limit request-per-menit.
            # Kalau error-nya nyebut "retry_delay" / detik tertentu, pakai itu; kalau tidak, pakai default.
            jeda = 45
            match_delay = re.search(r'retry_delay\s*\{\s*seconds:\s*(\d+)', error_text)
            if match_delay:
                jeda = int(match_delay.group(1)) + 5  # tambah buffer 5 detik
            print(f"[INFO] Menunggu {jeda} detik sebelum retry...")
            time.sleep(jeda)

# --- ROUTES AUTH ---
@app.route('/register', methods=['GET', 'POST'])
def register():
    if 'user_id' in session:
        return redirect(url_for('index'))

    if request.method == 'POST':
        nama = request.form.get('nama', '').strip()
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        konfirmasi = request.form.get('konfirmasi_password', '')

        if not nama or not email or not password:
            flash('Semua field wajib diisi!', 'danger')
            return redirect(url_for('register'))

        if password != konfirmasi:
            flash('Konfirmasi password tidak cocok!', 'danger')
            return redirect(url_for('register'))

        if len(password) < 6:
            flash('Password minimal 6 karakter!', 'danger')
            return redirect(url_for('register'))

        conn = get_db()
        existing = conn.execute('SELECT id FROM users WHERE email = ?', (email,)).fetchone()
        if existing:
            conn.close()
            flash('Email ini sudah terdaftar. Silakan login.', 'danger')
            return redirect(url_for('register'))

        password_hash = generate_password_hash(password)
        conn.execute('INSERT INTO users (nama, email, password_hash) VALUES (?, ?, ?)',
                     (nama, email, password_hash))
        conn.commit()
        conn.close()

        flash('Pendaftaran berhasil! Silakan login.', 'success')
        return redirect(url_for('login'))

    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('index'))

    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')

        conn = get_db()
        user = conn.execute('SELECT * FROM users WHERE email = ?', (email,)).fetchone()
        conn.close()

        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['user_nama'] = user['nama']
            flash(f"Selamat datang kembali, {user['nama']}!", 'success')
            return redirect(url_for('index'))
        else:
            flash('Email atau password salah.', 'danger')
            return redirect(url_for('login'))

    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    flash('Kamu sudah logout.', 'success')
    return redirect(url_for('login'))

# --- ROUTES ---
@app.route('/')
@login_required
def index():
    user_id = session['user_id']
    conn = get_db()
    total_materi = conn.execute('SELECT COUNT(*) FROM materi WHERE user_id = ?', (user_id,)).fetchone()[0]
    avg_nilai = conn.execute('''SELECT AVG(p.nilai) FROM progress p
                                 JOIN materi m ON p.materi_id = m.id
                                 WHERE m.user_id = ?''', (user_id,)).fetchone()[0] or 0
    total_flashcard = conn.execute('''SELECT COUNT(*) FROM hasil_belajar h
                                       JOIN materi m ON h.materi_id = m.id
                                       WHERE h.jenis_hasil="flashcard" AND m.user_id = ?''', (user_id,)).fetchone()[0]
    perlu_diulang = conn.execute('''SELECT COUNT(DISTINCT p.materi_id) FROM progress p
                                     JOIN materi m ON p.materi_id = m.id
                                     WHERE p.nilai < 70 AND m.user_id = ?''', (user_id,)).fetchone()[0]
    materi_terakhir = conn.execute('SELECT * FROM materi WHERE user_id = ? ORDER BY tanggal DESC LIMIT 4', (user_id,)).fetchall()
    conn.close()
    return render_template('index.html', total_materi=total_materi, avg_nilai=round(avg_nilai, 1),
                         total_flashcard=total_flashcard, perlu_diulang=perlu_diulang, materi_terakhir=materi_terakhir)

@app.route('/dashboard')
@login_required
def dashboard():
    user_id = session['user_id']
    conn = get_db()
    total_materi = conn.execute('SELECT COUNT(*) FROM materi WHERE user_id = ?', (user_id,)).fetchone()[0]
    avg_nilai = conn.execute('''SELECT AVG(p.nilai) FROM progress p
                                 JOIN materi m ON p.materi_id = m.id
                                 WHERE m.user_id = ?''', (user_id,)).fetchone()[0] or 0
    total_flashcard = conn.execute('''SELECT COUNT(*) FROM hasil_belajar h
                                       JOIN materi m ON h.materi_id = m.id
                                       WHERE h.jenis_hasil="flashcard" AND m.user_id = ?''', (user_id,)).fetchone()[0]
    perlu_diulang = conn.execute('''SELECT COUNT(DISTINCT p.materi_id) FROM progress p
                                     JOIN materi m ON p.materi_id = m.id
                                     WHERE p.nilai < 70 AND m.user_id = ?''', (user_id,)).fetchone()[0]
    materi_terakhir = conn.execute('SELECT * FROM materi WHERE user_id = ? ORDER BY tanggal DESC', (user_id,)).fetchall()

    # --- Data untuk Grafik Nilai Kuis ---
    riwayat_rows = conn.execute('''SELECT p.waktu_kuis, p.nilai, m.judul FROM progress p
                                    JOIN materi m ON p.materi_id = m.id
                                    WHERE m.user_id = ?
                                    ORDER BY p.waktu_kuis ASC''', (user_id,)).fetchall()
    riwayat_nilai = [{'tanggal': r['waktu_kuis'], 'nilai': r['nilai'], 'judul': r['judul']} for r in riwayat_rows]

    # --- Data untuk Streak & Heatmap Aktivitas (70 hari terakhir) ---
    aktivitas_rows = conn.execute('''
        SELECT hari, SUM(jumlah) as total FROM (
            SELECT DATE(tanggal) as hari, COUNT(*) as jumlah FROM materi WHERE user_id = ? GROUP BY hari
            UNION ALL
            SELECT DATE(p.waktu_kuis) as hari, COUNT(*) as jumlah FROM progress p
                JOIN materi m ON p.materi_id = m.id WHERE m.user_id = ? GROUP BY hari
        ) GROUP BY hari
    ''', (user_id, user_id)).fetchall()
    aktivitas_dict = {r['hari']: r['total'] for r in aktivitas_rows if r['hari']}

    max_nilai = conn.execute('''SELECT MAX(p.nilai) FROM progress p
                                 JOIN materi m ON p.materi_id = m.id
                                 WHERE m.user_id = ?''', (user_id,)).fetchone()[0] or 0
    jumlah_kuis_selesai = conn.execute('''SELECT COUNT(*) FROM progress p
                                           JOIN materi m ON p.materi_id = m.id
                                           WHERE m.user_id = ?''', (user_id,)).fetchone()[0]

    conn.close()

    # Hitung streak (hari beruntun aktif, dihitung mundur dari hari ini)
    hari_ini = date.today()
    streak = 0
    cek = hari_ini
    while aktivitas_dict.get(cek.isoformat(), 0) > 0:
        streak += 1
        cek -= timedelta(days=1)

    # --- Badge / Pencapaian ---
    badges = [
        {'nama': 'Kuis Pertama', 'icon': 'fa-flag-checkered', 'earned': jumlah_kuis_selesai >= 1},
        {'nama': 'Nilai Sempurna', 'icon': 'fa-trophy', 'earned': max_nilai >= 100},
        {'nama': '5 Materi Dibuat', 'icon': 'fa-layer-group', 'earned': total_materi >= 5},
        {'nama': '10 Materi Dibuat', 'icon': 'fa-crown', 'earned': total_materi >= 10},
        {'nama': 'Streak 3 Hari', 'icon': 'fa-fire', 'earned': streak >= 3},
        {'nama': 'Streak 7 Hari', 'icon': 'fa-fire-flame-curved', 'earned': streak >= 7},
        {'nama': 'Rata-rata di Atas 80', 'icon': 'fa-star', 'earned': avg_nilai >= 80 and jumlah_kuis_selesai >= 1},
    ]

    # Bangun grid 70 hari terakhir untuk heatmap
    aktivitas_grid = []
    for i in range(69, -1, -1):
        d = hari_ini - timedelta(days=i)
        jumlah = aktivitas_dict.get(d.isoformat(), 0)
        if jumlah == 0:
            level = 0
        elif jumlah <= 1:
            level = 1
        elif jumlah <= 2:
            level = 2
        elif jumlah <= 4:
            level = 3
        else:
            level = 4
        aktivitas_grid.append({'tanggal': d.strftime('%d %b %Y'), 'jumlah': jumlah, 'level': level})

    return render_template('dashboard.html', total_materi=total_materi, avg_nilai=round(avg_nilai, 1),
                         total_flashcard=total_flashcard, perlu_diulang=perlu_diulang, materi_terakhir=materi_terakhir,
                         riwayat_nilai=riwayat_nilai, aktivitas_grid=aktivitas_grid, streak=streak, badges=badges)

@app.route('/export_pdf/<int:id>')
@login_required
def export_pdf(id):
    conn = get_db()
    materi = conn.execute('SELECT * FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    if not materi:
        conn.close()
        flash('Materi tidak ditemukan.', 'danger')
        return redirect(url_for('dashboard'))

    hasil_rows = conn.execute('SELECT jenis_hasil, konten FROM hasil_belajar WHERE materi_id = ?', (id,)).fetchall()
    kuis_rows = conn.execute('SELECT soal, pilihan, kunci_jawaban, pembahasan FROM kuis WHERE materi_id = ?', (id,)).fetchall()
    conn.close()

    hasil_raw = {row['jenis_hasil']: row['konten'] for row in hasil_rows}

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(materi['judul'], styles['Title']))
    story.append(Paragraph(f"Jenjang: {materi['jenjang']} | Dibuat: {materi['tanggal']}", styles['Normal']))
    story.append(Spacer(1, 20))

    def tambah_section(judul, markdown_text):
        story.append(Paragraph(judul, styles['Heading1']))
        story.extend(markdown_ke_flowables(markdown_text, styles))
        story.append(Spacer(1, 16))

    tambah_section('Ringkasan Materi', hasil_raw.get('ringkasan', '-'))
    tambah_section('Poin-Poin Penting', hasil_raw.get('poin_penting', '-'))
    tambah_section('Kesimpulan', hasil_raw.get('kesimpulan', '-'))
    tambah_section('Rencana Belajar', hasil_raw.get('rencana_belajar', '-'))

    # Flashcard
    story.append(Paragraph('Flashcard', styles['Heading1']))
    try:
        flashcards = json.loads(hasil_raw.get('flashcard', '[]'))
    except Exception:
        flashcards = []
    if not flashcards:
        story.append(Paragraph('-', styles['Normal']))
    for i, fc in enumerate(flashcards, 1):
        story.append(Paragraph(f"<b>{i}. {fc.get('q', '')}</b>", styles['Normal']))
        story.append(Paragraph(f"Jawaban: {fc.get('a', '')}", styles['Normal']))
        story.append(Spacer(1, 8))
    story.append(Spacer(1, 16))

    # Kuis
    story.append(Paragraph('Kuis', styles['Heading1']))
    if not kuis_rows:
        story.append(Paragraph('-', styles['Normal']))
    for i, soal in enumerate(kuis_rows, 1):
        story.append(Paragraph(f"<b>{i}. {soal['soal']}</b>", styles['Normal']))
        try:
            opsi = json.loads(soal['pilihan'])
        except Exception:
            opsi = []
        for op in opsi:
            story.append(Paragraph(op, styles['Normal']))
        story.append(Paragraph(f"<i>Kunci jawaban: {soal['kunci_jawaban']} &mdash; {soal['pembahasan']}</i>", styles['Normal']))
        story.append(Spacer(1, 10))

    doc.build(story)
    buffer.seek(0)

    nama_file = re.sub(r'[^a-zA-Z0-9_-]', '_', materi['judul']) + '.pdf'
    return send_file(buffer, as_attachment=True, download_name=nama_file, mimetype='application/pdf')

@app.route('/buat_materi', methods=['GET', 'POST'])
@login_required
def buat_materi():
    if request.method == 'POST':
        judul = request.form['judul']
        jenjang = request.form['jenjang']
        tipe_input = request.form['tipe_input']

        teks = ""
        if 'file_materi' in request.files:
            file = request.files['file_materi']
            if file.filename != '':
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
                file.save(filepath)
                teks = extract_text_from_file(filepath, file.filename)
                os.remove(filepath)

        if not teks and 'teks_langsung' in request.form:
            teks = request.form['teks_langsung']

        if not teks.strip():
            flash('Materi tidak boleh kosong!', 'danger')
            return redirect(url_for('buat_materi'))

        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('INSERT INTO materi (judul, jenjang, tipe_input, isi_teks, user_id) VALUES (?, ?, ?, ?, ?)',
                       (judul, jenjang, tipe_input, teks, session['user_id']))
        materi_id = cursor.lastrowid

        flash('Membuat paket belajar... Mohon tunggu (bisa memakan waktu 10-20 detik).', 'info')
        conn.commit()
        conn.close()

        # Generate AI
        hasil = generate_paket_belajar(teks, jenjang)

        conn = get_db()
        cursor = conn.cursor()
        for jenis, konten in hasil.items():
            cursor.execute('INSERT INTO hasil_belajar (materi_id, jenis_hasil, konten) VALUES (?, ?, ?)',
                           (materi_id, jenis, konten))

        # Simpan Kuis
        if 'kuis' in hasil:
            try:
                kuis_data = clean_json(hasil['kuis']) if isinstance(hasil['kuis'], str) else hasil['kuis']
                for soal in kuis_data:
                    cursor.execute('INSERT INTO kuis (materi_id, soal, pilihan, kunci_jawaban, pembahasan) VALUES (?, ?, ?, ?, ?)',
                                   (materi_id, soal['soal'], json.dumps(soal.get('opsi', [])), soal.get('kunci', ''), soal.get('pembahasan', '')))
            except Exception as e:
                print(f"Error simpan kuis: {e}")

        conn.commit()
        conn.close()

        flash('Paket belajar berhasil dibuat!', 'success')
        return redirect(url_for('materi_detail', id=materi_id))

    return render_template('buat_materi.html')

@app.route('/materi/<int:id>')
@login_required
def materi_detail(id):
    conn = get_db()
    materi = conn.execute('SELECT * FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    if not materi:
        conn.close()
        flash('Materi tidak ditemukan.', 'danger')
        return redirect(url_for('dashboard'))
    hasil_rows = conn.execute('SELECT jenis_hasil, konten FROM hasil_belajar WHERE materi_id = ?', (id,)).fetchall()
    kuis_count = conn.execute('SELECT COUNT(*) FROM kuis WHERE materi_id = ?', (id,)).fetchone()[0]
    progress = conn.execute('SELECT * FROM progress WHERE materi_id = ? ORDER BY waktu_kuis DESC LIMIT 1', (id,)).fetchone()
    conn.close()

    # Field yang isinya teks/markdown biasa (dirender jadi HTML)
    field_markdown = {'ringkasan', 'poin_penting', 'kesimpulan', 'rencana_belajar'}

    hasil_dict = {}
    flashcard_json = '[]'
    for row in hasil_rows:
        jenis = row['jenis_hasil']
        konten = row['konten']
        if jenis in field_markdown:
            hasil_dict[jenis] = markdown.markdown(konten)
        elif jenis == 'flashcard':
            flashcard_json = konten  # simpan mentah (JSON string), JANGAN di-markdown
        else:
            hasil_dict[jenis] = konten

    return render_template('materi_detail.html', materi=materi, hasil=hasil_dict, kuis_count=kuis_count, progress=progress, flashcard_json=flashcard_json)

@app.route('/jelaskan/<int:id>', methods=['POST'])
@login_required
def jelaskan(id):
    """Endpoint AJAX untuk fitur 'Jelaskan Lebih Mudah'. Menerima gaya penjelasan,
    memanggil AI dengan prompt sesuai gaya, dan mengembalikan hasilnya sebagai JSON."""
    conn = get_db()
    materi = conn.execute('SELECT * FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    conn.close()

    if not materi:
        return jsonify({'sukses': False, 'pesan': 'Materi tidak ditemukan.'}), 404

    data = request.get_json(silent=True) or {}
    gaya = data.get('gaya', 'sd')

    instruksi_gaya = {
        'sd': "Jelaskan ulang materi ini dengan bahasa yang SANGAT SEDERHANA, seolah menjelaskan ke anak SD kelas 4-5. Gunakan kalimat pendek dan kata-kata sehari-hari.",
        'analogi': "Jelaskan ulang materi ini dengan MENGGUNAKAN ANALOGI/PERUMPAMAAN dari hal-hal yang akrab dalam kehidupan sehari-hari, supaya konsepnya lebih mudah dibayangkan.",
        'contoh': "Jelaskan ulang materi ini dengan memberikan CONTOH-CONTOH NYATA dan konkret dari kehidupan sehari-hari untuk tiap konsep pentingnya.",
    }
    instruksi = instruksi_gaya.get(gaya, instruksi_gaya['sd'])

    prompt = f"""{instruksi}

Gunakan format Markdown (heading/bullet points jika perlu). Jawaban maksimal sekitar 3-4 paragraf, jangan terlalu panjang.

MATERI ASLI:
{materi['isi_teks'][:3000]}
"""

    try:
        response = model.generate_content(prompt, request_options={"timeout": 60})
        penjelasan_html = markdown.markdown(response.text)
        return jsonify({'sukses': True, 'penjelasan': penjelasan_html})
    except Exception as e:
        error_text = str(e)
        print(f"[ERROR] Gagal generate penjelasan (gaya={gaya}): {error_text}")
        if "PerDay" in error_text:
            pesan = "Kuota harian Gemini API sudah habis. Coba lagi nanti setelah reset."
        else:
            pesan = "Gagal membuat penjelasan. Coba lagi sebentar lagi."
        return jsonify({'sukses': False, 'pesan': pesan}), 500

@app.route('/kuis/<int:id>', methods=['GET', 'POST'])
@login_required
def kuis(id):
    conn = get_db()
    materi = conn.execute('SELECT * FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    if not materi:
        conn.close()
        flash('Materi tidak ditemukan.', 'danger')
        return redirect(url_for('dashboard'))

    if request.method == 'POST':
        soal_list = conn.execute('SELECT * FROM kuis WHERE materi_id = ?', (id,)).fetchall()
        jawaban = dict(request.form)
        benar = 0
        total = len(soal_list)

        for soal in soal_list:
            if str(soal['id']) in jawaban and jawaban[str(soal['id'])] == soal['kunci_jawaban']:
                benar += 1

        nilai = (benar / total * 100) if total > 0 else 0
        conn.execute('INSERT INTO progress (materi_id, total_soal, benar, nilai, waktu_kuis) VALUES (?, ?, ?, ?, ?)',
                     (id, total, benar, nilai, datetime.now()))
        conn.commit()

        # Parse pilihan untuk template
        soal_list = [dict(soal) for soal in soal_list]
        for soal in soal_list:
            try: soal['pilihan'] = json.loads(soal['pilihan'])
            except: soal['pilihan'] = ['A. Opsi 1', 'B. Opsi 2', 'C. Opsi 3', 'D. Opsi 4']

        conn.close()
        return render_template('hasil_kuis.html', materi=materi, benar=benar, total=total, nilai=round(nilai, 1), soal_list=soal_list, jawaban=jawaban)

    soal_list = conn.execute('SELECT * FROM kuis WHERE materi_id = ?', (id,)).fetchall()
    soal_list = [dict(soal) for soal in soal_list]
    for soal in soal_list:
        try: soal['pilihan'] = json.loads(soal['pilihan'])
        except: soal['pilihan'] = ['A. Opsi 1', 'B. Opsi 2', 'C. Opsi 3', 'D. Opsi 4']
    conn.close()
    return render_template('kuis.html', materi=materi, soal_list=soal_list)

@app.route('/hapus_materi/<int:id>', methods=['POST'])
@login_required
def hapus_materi(id):
    conn = get_db()
    materi = conn.execute('SELECT id FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    if not materi:
        conn.close()
        flash('Materi tidak ditemukan.', 'danger')
        return redirect(url_for('dashboard'))
    conn.execute('DELETE FROM materi WHERE id = ?', (id,))
    conn.execute('DELETE FROM hasil_belajar WHERE materi_id = ?', (id,))
    conn.execute('DELETE FROM kuis WHERE materi_id = ?', (id,))
    conn.execute('DELETE FROM progress WHERE materi_id = ?', (id,))
    conn.commit()
    conn.close()
    flash('Materi berhasil dihapus!', 'success')
    return redirect(url_for('dashboard'))

@app.route('/belajar/<int:id>')
@login_required
def belajar(id):
    conn = get_db()
    materi = conn.execute('SELECT * FROM materi WHERE id = ? AND user_id = ?', (id, session['user_id'])).fetchone()
    if not materi:
        conn.close()
        flash('Materi tidak ditemukan.', 'danger')
        return redirect(url_for('dashboard'))
    flashcard_raw = conn.execute('SELECT konten FROM hasil_belajar WHERE materi_id = ? AND jenis_hasil = "flashcard"', (id,)).fetchone()

    flashcards = []
    if flashcard_raw:
        try: flashcards = json.loads(flashcard_raw['konten'])
        except: pass
    conn.close()
    return render_template('belajar.html', materi=materi, flashcards=flashcards)

# Initialize
init_db()

if __name__ == '__main__':
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    app.run(debug=True, port=5000) # Pakai port 5000 biar aman dari bentrok XAMPP